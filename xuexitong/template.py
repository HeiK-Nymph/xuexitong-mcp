"""DOCX 读取与"空白位置提示"。

定位变化（重要）：MCP **不再预设待填字段**。它只做两件事：
  1. 读取文档全文（交给当前窗口自己阅读、判断、作答）
  2. 顺带给出"这里看起来是空白/待填"的**提示**（hints），仅作参考，模型可忽略

个人信息识别保留：姓名/学号/班级/日期 这类字段任何情况下都不得由模型编造。
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import List, Optional

PLACEHOLDER_RE = re.compile(r"\{\{(\w+)\}\}")
UNDERSCORE_PATTERN = re.compile(r"^[＿_]+$")
BLANK_RE = re.compile(r"^\s*$")
SECTION_TITLE_RE = re.compile(r"^[一二三四五六七八九十]+[、.．]")
COLON_END_RE = re.compile(r"：$")
INLINE_BLANK_RE = re.compile(r"[：:][ \u3000\t]{4,}")
QUESTION_END_RE = re.compile(r"[?？]\s*$")
STEP_KEYWORDS = ("定义", "创建", "编写", "启动", "提交", "实现",
                 "设计", "构建", "搭建", "配置", "测试", "运行")
SCREENSHOT_PATTERNS = [
    re.compile(r"^[（(]截图[)）]$"), re.compile(r"^[（(]运行结果[)）]$"),
    re.compile(r"^[（(]实验截图[)）]$"), re.compile(r"^[（(]运行结果截图[)）]$"),
    re.compile(r"^[（(]实验[)）].*截图"), re.compile(r"^[（(]可附图[)）]"),
]

# 个人信息 / 无法从附件推断的字段：模型不得编造，只能由用户提供
ASK_USER_KEYWORDS = ("姓名", "学号", "班级", "学院", "系别", "专业", "组号", "同组",
                     "指导教师", "指导老师", "任课教师", "老师", "签名", "成绩",
                     "日期", "时间", "地点", "台号", "联系方式", "电话")


@dataclass
class BlankHint:
    """"这里看起来是空的" —— 仅提示，不是必须填写的字段。"""

    paragraph_index: int
    kind: str                 # blank | inline | underscore | screenshot | placeholder
    after_text: str = ""      # 上一段（通常是要点标签）
    hint_text: str = ""       # 该位置原文（inline/underscore/screenshot 用）
    paragraph_count: int = 1
    personal_fields: List[str] = field(default_factory=list)

    @property
    def is_personal(self) -> bool:
        return bool(self.personal_fields)


@dataclass
class DocumentText:
    path: str
    paragraphs: List[str] = field(default_factory=list)
    hints: List[BlankHint] = field(default_factory=list)
    error: str = ""

    @property
    def text(self) -> str:
        return "\n".join(self.paragraphs)


def personal_fields_in(text: str) -> List[str]:
    """挑出文本里出现的个人信息字段名（去重、保持顺序）。"""
    found: List[str] = []
    for kw in ASK_USER_KEYWORDS:
        if kw in (text or "") and kw not in found:
            found.append(kw)
    return found


def read_document(path: str | Path) -> DocumentText:
    """读 .docx 全文 + 空白位置提示。非 .docx 会报错（.doc 需先转换）。"""
    p = Path(path)
    if p.suffix.lower() != ".docx":
        return DocumentText(path=str(p), error=f"只支持 .docx（当前 {p.suffix}），请先转换")
    try:
        from docx import Document as DocxDocument

        paragraphs = [para.text for para in DocxDocument(str(p)).paragraphs]
    except Exception as e:  # noqa: BLE001
        return DocumentText(path=str(p), error=f"打开失败：{e}")
    return DocumentText(path=str(p), paragraphs=paragraphs, hints=detect_hints(paragraphs))


def detect_hints(paragraphs: List[str]) -> List[BlankHint]:
    """找"看起来可以填空"的位置（供模型参考，不作强制）。"""
    hints: List[BlankHint] = []
    blanks = [BLANK_RE.match(p or "") is not None for p in paragraphs]

    for i, raw in enumerate(paragraphs):
        text = (raw or "").strip()

        m = PLACEHOLDER_RE.search(text)
        if m:
            hints.append(BlankHint(i, "placeholder", hint_text=text))
            continue
        if any(p.match(text) for p in SCREENSHOT_PATTERNS):
            hints.append(BlankHint(i, "screenshot", hint_text=text,
                                   after_text=paragraphs[i - 1].strip() if i else ""))
            continue
        if UNDERSCORE_PATTERN.match(text):
            before = paragraphs[i - 1].strip() if i else ""
            hints.append(BlankHint(i, "underscore", after_text=before, hint_text=text,
                                   personal_fields=personal_fields_in(before)))
            continue
        if len(text) <= 80 and INLINE_BLANK_RE.search(text):
            hints.append(BlankHint(i, "inline", hint_text=raw.strip(),
                                   personal_fields=personal_fields_in(text)))
            continue

        question_like = QUESTION_END_RE.search(text) is not None
        step_like = COLON_END_RE.search(text) is not None and len(text) <= 40 or any(k in text for k in STEP_KEYWORDS)
        chapter_like = SECTION_TITLE_RE.match(text) is not None
        if (question_like or step_like or chapter_like) and i + 1 < len(paragraphs) and blanks[i + 1]:
            count = 0
            j = i + 1
            while j < len(paragraphs) and blanks[j]:
                count += 1
                j += 1
            before = text[:120]
            hints.append(BlankHint(i + 1, "blank", after_text=before,
                                   paragraph_count=count,
                                   personal_fields=personal_fields_in(before)))
    return hints
