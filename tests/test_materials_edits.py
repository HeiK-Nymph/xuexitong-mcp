"""取材料 + 锚点写回 的离线测试。

钉住四条硬规则：
  · 压缩包要解开，并把包内文档内容交给窗口
  · 锚点写回：after / replace / append 三种模式，多行要插入新段落
  · **个人信息不得编造**：模型给的值一律拒绝，只认 user_fields
  · 信息不足（INSUFFICIENT）→ 标【待补】，不编内容
"""

from __future__ import annotations

import sys
import tempfile
import unittest
import zipfile
from pathlib import Path
from types import SimpleNamespace
from unittest import mock

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from xuexitong import tasks  # noqa: E402
from xuexitong.docx_fill import apply_edits  # noqa: E402
from xuexitong.template import read_document  # noqa: E402


def make_docx(path: Path, lines: list[str]) -> None:
    from docx import Document

    doc = Document()
    for line in lines:
        doc.add_paragraph(line)
    doc.save(str(path))


class ApplyEditsTest(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        self.tpl = self.tmp / "模板.docx"
        make_docx(self.tpl, [
            "实验三 XXX",
            "班级：                     姓名：            学号：",     # 个人信息行
            "1. 回答问题：为什么？",
            "",
            "三、实验小结：（存在问题、解决办法、实验心得）",
            "",
        ])

    def test_after_inserts_multi_line_paragraphs(self):
        edits = [{"anchor": "1. 回答问题：为什么？", "content": "第一行\n第二行\n第三行"}]
        rep = apply_edits(self.tpl, edits, out_path=self.tmp / "o1.docx")
        text = [p.text for p in __import__("docx").Document(rep.out_path).paragraphs]
        i = text.index("1. 回答问题：为什么？")
        self.assertEqual(text[i + 1: i + 4], ["第一行", "第二行", "第三行"])
        self.assertEqual(rep.applied[0]["mode"], "after")
        self.assertEqual(rep.not_found, [])

    def test_replace_and_append(self):
        edits = [
            {"anchor": "三、实验小结：（存在问题、解决办法、实验心得）", "content": "小结：一切正常", "mode": "replace"},
            {"anchor": "", "content": "（附加说明）", "mode": "append"},
        ]
        rep = apply_edits(self.tpl, edits, out_path=self.tmp / "o2.docx")
        text = [p.text for p in __import__("docx").Document(rep.out_path).paragraphs]
        self.assertIn("小结：一切正常", text)
        self.assertNotIn("三、实验小结：（存在问题、解决办法、实验心得）", text)
        self.assertEqual(text[-1], "（附加说明）")

    def test_anchor_not_found_is_reported(self):
        rep = apply_edits(self.tpl, [{"anchor": "根本不存在的段落", "content": "x"}],
                          out_path=self.tmp / "o3.docx")
        self.assertEqual(rep.applied, [])
        self.assertEqual(rep.not_found, ["根本不存在的段落"])

    def test_personal_info_is_never_filled_by_model(self):
        """核心规则：模型想填姓名/学号/班级 → 必须拒绝。"""
        edits = [{"anchor": "班级：                     姓名：            学号：",
                  "content": "班级：计科241  姓名：虚拟名字  学号：0000000000",
                  "mode": "replace"}]
        rep = apply_edits(self.tpl, edits, out_path=self.tmp / "o4.docx")
        text = [p.text for p in __import__("docx").Document(rep.out_path).paragraphs]
        joined = "\n".join(text)
        self.assertNotIn("虚拟名字", joined)          # 编造的值绝不允许落盘
        self.assertIn("＿＿＿（待填）", joined)
        self.assertEqual(len(rep.refused_personal), 1)
        self.assertTrue(rep.needs_user and rep.needs_user[0]["fields"])

    def test_user_provided_personal_info_is_written(self):
        edits = [{"anchor": "班级：                     姓名：            学号：",
                  "content": "（模型给的值应被忽略）", "mode": "replace"}]
        rep = apply_edits(self.tpl, edits, out_path=self.tmp / "o5.docx",
                          user_fields={"班级": "计科241", "姓名": "张三"})
        text = "\n".join(p.text for p in __import__("docx").Document(rep.out_path).paragraphs)
        self.assertIn("班级：计科241", text)
        self.assertIn("姓名：张三", text)
        self.assertIn("学号：＿＿＿（待填）", text)      # 用户没给学号 → 留待填
        self.assertEqual(rep.needs_user[0]["fields"], ["学号"])

    def test_insufficient_content_is_marked_not_invented(self):
        edits = [{"anchor": "1. 回答问题：为什么？",
                  "content": {"INSUFFICIENT": True, "reason": "被分析对象不在附件里，需要你提供原始数据"}}]
        rep = apply_edits(self.tpl, edits, out_path=self.tmp / "o6.docx")
        text = "\n".join(p.text for p in __import__("docx").Document(rep.out_path).paragraphs)
        self.assertIn("【待补：被分析对象不在附件里，需要你提供原始数据】", text)
        self.assertEqual(len(rep.insufficient), 1)
        self.assertEqual(rep.applied, [])


class GetMaterialsTest(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        self.homework = "测试作业"
        self.folder = self.tmp / "acct" / self.homework
        self.folder.mkdir(parents=True)
        self.fake_hw = SimpleNamespace(work_id="999", name=self.homework)

    def _call(self, **kw):
        end = self.folder
        self.folder = end
        with mock.patch.object(tasks, "_resolve_homework", return_value=(self.fake_hw, "测试课程")), \
             mock.patch.object(tasks, "get_client", return_value=object()), \
             mock.patch.object(tasks.store, "account_dir_name", return_value="acct"):
            return tasks.get_homework_materials(work_id="999", root=str(self.tmp), **kw)

    def test_no_attachment_is_need_more_info(self):
        res = self._call()
        self.assertEqual(res["status"], "NEED_MORE_INFO")
        self.assertIn("没有任何附件", res["reasons"][0])
        self.assertIn("不会编造", res["summary"])

    def test_only_image_is_need_more_info(self):
        (self.folder / "1.jpg").write_bytes(b"\xff\xd8\xff\xe0fake")
        res = self._call()
        self.assertEqual(res["status"], "NEED_MORE_INFO")
        self.assertIn("没有可读文档", res["reasons"][0])

    def test_zip_is_extracted_and_text_returned(self):
        inner = self.folder / "inner.docx"
        make_docx(inner, ["包内模板第一行", "班级：      姓名："])
        zpath = self.folder / "原始附件.zip"
        with zipfile.ZipFile(zpath, "w") as z:
            z.write(inner, "inner.docx")
        inner.unlink()

        res = self._call()
        self.assertEqual(res["status"], "OK", res.get("reasons"))
        self.assertTrue(res["extract_notes"] and "解压出 1 个文件" in res["extract_notes"][0])
        self.assertEqual(len(res["documents"]), 1)
        self.assertIn("包内模板第一行", res["documents"][0]["text"])
        self.assertIn("姓名", res["personal_fields"])
        self.assertIsNotNone(res["template_suggestion"])

    def test_docx_text_has_line_numbers_for_anchors(self):
        make_docx(self.folder / "t.docx", ["甲", "乙"])
        res = self._call()
        self.assertIn("   0| 甲", res["documents"][0]["text"])


class MultiAnchorSafetyTest(unittest.TestCase):
    """回归：前面的锚点插入多行后，后面的锚点仍要落在正确位置。

    旧实现按 paragraph_index 定位，插入段落会让后面索引整体漂移、内容写错位置；
    改成锚点定位后这类问题从设计上消失，仍用本条守住。
    """

    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        self.tpl = self.tmp / "multi.docx"
        make_docx(self.tpl, [
            "1.标题A：", "", "中间的原文，不能被覆盖", "2.标题B：", "", "结尾原文",
        ])

    def test_later_anchor_survives_previous_insertions(self):
        edits = [
            {"anchor": "1.标题A：", "content": "A1\nA2\nA3"},
            {"anchor": "2.标题B：", "content": "B的内容"},
            {"anchor": "结尾原文", "content": "尾部补充", "mode": "after"},
        ]
        out = self.tmp / "out.docx"
        rep = apply_edits(self.tpl, edits, out_path=out)
        self.assertEqual(len(rep.applied), 3)
        self.assertEqual(rep.not_found, [])

        from docx import Document
        text = [p.text for p in Document(str(out)).paragraphs]
        ia = text.index("1.标题A：")
        self.assertEqual(text[ia + 1: ia + 4], ["A1", "A2", "A3"])
        self.assertIn("中间的原文，不能被覆盖", text)
        ib = text.index("2.标题B：")
        self.assertEqual(text[ib + 1], "B的内容")
        ie = text.index("结尾原文")
        self.assertEqual(text[ie + 1], "尾部补充")


if __name__ == "__main__":
    unittest.main(verbosity=1)
